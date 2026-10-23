# -*- coding: utf-8 -*-
import os, re, subprocess, sys
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = Path(__file__).resolve().parent
ASSETS = BASE / "assets"
DOCX = BASE / "Blood_Bank_Donor_Management_Report.docx"
PDF = BASE / "Blood_Bank_Donor_Management_Report.pdf"
INK="1F2421"; SOFT="55594F"; FAINT="62655A"; RED="8C1D18"; DEEP="6B120F"
AMBER="D9A441"; TEAL="2F5D57"; PAPER="F3EDE1"; RAISED="FAF6EC"; LINE="CFC4B2"; STRONG="A99C86"
FONT="Arial"; MONO="Menlo"

doc = Document()
sec = doc.sections[0]
sec.page_width=Inches(8.27); sec.page_height=Inches(11.69)
sec.left_margin=Inches(0.82); sec.right_margin=Inches(0.82)
sec.top_margin=Inches(0.72); sec.bottom_margin=Inches(0.72)
sec.header_distance=Inches(0.34); sec.footer_distance=Inches(0.34)
sec.different_first_page_header_footer = True

# ---------- core styles ----------
def set_font(style, size, color=INK, bold=False, name=FONT):
    style.font.name=name; style.font.size=Pt(size); style.font.bold=bold
    style.font.color.rgb=RGBColor.from_string(color)
    rpr=style.element.get_or_add_rPr(); rf=rpr.find(qn('w:rFonts'))
    if rf is None:
        rf=OxmlElement('w:rFonts'); rpr.append(rf)
    rf.set(qn('w:ascii'),name); rf.set(qn('w:hAnsi'),name); rf.set(qn('w:cs'),name); rf.set(qn('w:eastAsia'),name)

normal=doc.styles['Normal']; set_font(normal,10.4,INK)
normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.25
normal.paragraph_format.widow_control=True

for name,size,color,bold,before,after in [('Heading 1',17,RED,True,16,7),('Heading 2',12.3,INK,True,13,4),('Heading 3',10.6,SOFT,True,10,3)]:
    st=doc.styles[name]; set_font(st,size,color,bold)
    st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after)
    st.paragraph_format.keep_with_next=True; st.paragraph_format.line_spacing=1.12

for name in ['List Bullet','List Number']:
    st=doc.styles[name]; set_font(st,10.4,INK)
    st.paragraph_format.space_after=Pt(3); st.paragraph_format.line_spacing=1.22
    st.paragraph_format.left_indent=Inches(0.26); st.paragraph_format.first_line_indent=Inches(-0.16)

def shade(cell, hexcolor):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd')
    shd.set(qn('w:val'),'clear'); shd.set(qn('w:color'),'auto'); shd.set(qn('w:fill'),hexcolor); tcPr.append(shd)

def cell_margins(cell, top=70, left=95, bottom=70, right=95):
    tcPr=cell._tc.get_or_add_tcPr(); mar=OxmlElement('w:tcMar')
    for side,val in [('top',top),('left',left),('bottom',bottom),('right',right)]:
        e=OxmlElement(f'w:{side}'); e.set(qn('w:w'),str(val)); e.set(qn('w:type'),'dxa'); mar.append(e)
    tcPr.append(mar)

def set_borders(table, color=LINE, sz=4, inside=True):
    tblPr=table._tbl.tblPr; borders=OxmlElement('w:tblBorders')
    edges=['top','left','bottom','right'] + (['insideH','insideV'] if inside else [])
    for edge in edges:
        e=OxmlElement(f'w:{edge}'); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),str(sz))
        e.set(qn('w:space'),'0'); e.set(qn('w:color'),color); borders.append(e)
    tblPr.append(borders)

def no_borders(table):
    tblPr=table._tbl.tblPr; borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement(f'w:{edge}'); e.set(qn('w:val'),'nil'); borders.append(e)
    tblPr.append(borders)

def row_no_split(row):
    trPr=row._tr.get_or_add_trPr(); e=OxmlElement('w:cantSplit'); trPr.append(e)

def repeat_header(row):
    trPr=row._tr.get_or_add_trPr(); e=OxmlElement('w:tblHeader'); e.set(qn('w:val'),'true'); trPr.append(e)

def para_bottom_border(p, color=LINE, sz=6, space=4):
    pPr=p._p.get_or_add_pPr(); pbdr=OxmlElement('w:pBdr'); b=OxmlElement('w:bottom')
    b.set(qn('w:val'),'single'); b.set(qn('w:sz'),str(sz)); b.set(qn('w:space'),str(space)); b.set(qn('w:color'),color)
    pbdr.append(b); pPr.append(pbdr)

def add_run(p, text, size=None, color=None, bold=False, italic=False, name=FONT, caps=False):
    r=p.add_run(text); r.font.name=name; r.bold=bold; r.italic=italic
    if size: r.font.size=Pt(size)
    if color: r.font.color.rgb=RGBColor.from_string(color)
    if caps: r.font.all_caps=True
    return r

def body(text, lead=None, size=10.4, after=6, color=INK, align=None):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(after)
    if align is not None: p.alignment=align
    if lead: add_run(p,lead,size=size,color=color,bold=True)
    add_run(p,text,size=size,color=color)
    return p

def bullet(text, bold_lead=None):
    p=doc.add_paragraph(style='List Bullet')
    if bold_lead: add_run(p,bold_lead,size=10.4,bold=True)
    add_run(p,text,size=10.4)
    return p

def caption(text):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(3); p.paragraph_format.space_after=Pt(11)
    p.paragraph_format.keep_together=True
    add_run(p,text,size=8.4,color=FAINT,italic=True)

def figure(path, text, width=6.4):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(4); p.paragraph_format.space_after=Pt(2)
    p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(path), width=Inches(width))
    caption(text)

def table(headers, rows, widths, font=8.5, header_font=8.2, align=None, first_col_mono=False):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False; set_borders(t)
    hdr=t.rows[0]; repeat_header(hdr); row_no_split(hdr)
    for i,h in enumerate(headers):
        c=hdr.cells[i]; c.width=Inches(widths[i]); shade(c,RED); cell_margins(c)
        p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.05
        add_run(p,str(h),size=header_font,color="FFFFFF",bold=True)
    for ridx,row in enumerate(rows):
        cells=t.add_row().cells; row_no_split(t.rows[-1])
        for i,val in enumerate(row):
            c=cells[i]; c.width=Inches(widths[i]); cell_margins(c)
            if ridx%2==1: shade(c,PAPER)
            p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.12
            if align and i in align: p.alignment=align[i]
            txt=str(val)
            if first_col_mono and i==0:
                add_run(p,txt,size=font-0.3,color=DEEP,bold=True,name=MONO)
            else:
                add_run(p,txt,size=font)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

def callout(title, text):
    t=doc.add_table(rows=1, cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    set_borders(t, color=LINE, sz=4, inside=False)
    c=t.rows[0].cells[0]; c.width=Inches(6.55); shade(c,PAPER); cell_margins(c,90,130,90,130)
    p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(3)
    add_run(p,title,size=8.4,color=RED,bold=True,caps=True)
    p2=c.add_paragraph(); p2.paragraph_format.space_after=Pt(0); p2.paragraph_format.line_spacing=1.2
    add_run(p2,text,size=9.2,color=SOFT)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

def code_block(lines):
    p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(0.12); p.paragraph_format.right_indent=Inches(0.12)
    p.paragraph_format.space_before=Pt(5); p.paragraph_format.space_after=Pt(10); p.paragraph_format.keep_together=True
    p.paragraph_format.line_spacing=1.12
    pPr=p._p.get_or_add_pPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:val'),'clear'); shd.set(qn('w:color'),'auto'); shd.set(qn('w:fill'),"F6F1E6"); pPr.append(shd)
    for i,line in enumerate(lines):
        if i: p.add_run().add_break()
        add_run(p,line,size=8.1,color=INK,name=MONO)

def field_run(p, instr):
    r=p.add_run(); fld=OxmlElement('w:fldChar'); fld.set(qn('w:fldCharType'),'begin'); r._r.append(fld)
    r2=p.add_run(); it=OxmlElement('w:instrText'); it.set(qn('xml:space'),'preserve'); it.text=instr; r2._r.append(it)
    r3=p.add_run(); f2=OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'),'separate'); r3._r.append(f2)
    r4=p.add_run('1')
    r5=p.add_run(); f3=OxmlElement('w:fldChar'); f3.set(qn('w:fldCharType'),'end'); r5._r.append(f3)

# ---------- header / footer ----------
hdr=sec.header
hp=hdr.paragraphs[0]; hp.paragraph_format.tab_stops.add_tab_stop(Inches(6.63),WD_TAB_ALIGNMENT.RIGHT)
hp.paragraph_format.space_after=Pt(2)
add_run(hp,"Blood Bank & Donor Management System",size=8,color=FAINT)
add_run(hp,"\tPBL Implementation Report · 2026–27",size=8,color=FAINT)
para_bottom_border(hp,LINE,4,3)

ftr=sec.footer
fp=ftr.paragraphs[0]; fp.paragraph_format.tab_stops.add_tab_stop(Inches(6.63),WD_TAB_ALIGNMENT.RIGHT)
add_run(fp,"Department of Computer Science  ·  Academic prototype",size=7.8,color=FAINT)
add_run(fp,"\tPage ",size=7.8,color=FAINT); field_run(fp,"PAGE"); add_run(fp," of ",size=7.8,color=FAINT); field_run(fp,"NUMPAGES")

# ---------- cover ----------
sp=doc.add_paragraph(); sp.paragraph_format.space_after=Pt(18)
band=doc.add_table(rows=1, cols=1); band.autofit=False; no_borders(band)
c=band.rows[0].cells[0]; c.width=Inches(6.63); shade(c,DEEP); cell_margins(c,120,150,120,150)
p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0)
add_run(p,"PROJECT-BASED LEARNING (PBL) REPORT   ·   IMPLEMENTATION EDITION",size=8.3,color="F5E6D8",bold=True)

p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(26); p.paragraph_format.space_after=Pt(5)
add_run(p,"ACADEMIC YEAR 2026–27   ·   HEALTHCARE INFORMATION SYSTEMS",size=8.2,color=RED,bold=True)
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(9); p.paragraph_format.line_spacing=0.98
add_run(p,"Online Blood Bank and\nDonor Management System",size=31,color=INK,bold=True)
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(8); p.paragraph_format.line_spacing=1.16
add_run(p,"A relational database application with role-based dashboards, live stock, request approval and donation records.",size=11.5,color=SOFT)
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(20); para_bottom_border(p,RED,10,1)

meta=[("Student name","[Your Name]"),("Roll number","[Roll No.]"),("Class / division","[Class & Division]"),
      ("Department","[Department]"),("Institution","[Institution]"),("Faculty guide","[Guide Name]"),
      ("Project domain","Healthcare Information Systems"),("Technology","Python · Flask · SQLite · HTML/CSS/JS"),
      ("Submission date","[DD Month 2026]")]
t=doc.add_table(rows=0, cols=2); t.autofit=False; no_borders(t)
for k,v in meta:
    cells=t.add_row().cells
    cells[0].width=Inches(1.55); cells[1].width=Inches(5.08)
    for cc in cells: cell_margins(cc,45,0,45,0)
    p=cells[0].paragraphs[0]; p.paragraph_format.space_after=Pt(0)
    add_run(p,k.upper(),size=7.6,color=FAINT,bold=True)
    p=cells[1].paragraphs[0]; p.paragraph_format.space_after=Pt(0)
    add_run(p,v,size=9.3,color=INK)

p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(38); p.paragraph_format.space_after=Pt(3)
add_run(p,"Submitted in partial fulfilment of the Project-Based Learning (PBL) component.",size=8.8,color=SOFT,italic=True)
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(0)
add_run(p,"All names, hospitals, stock counts and requests in this report are synthetic demonstration data. This prototype is not a clinical system.",size=8.2,color=FAINT)
doc.add_page_break()

# ---------- abstract ----------
h=doc.add_heading("Abstract",level=1)
body("Blood is a time-critical clinical resource, yet many small blood banks still coordinate donors, stock and hospital requests through disconnected registers, spreadsheets and phone calls. That fragmentation makes it difficult to answer a simple question quickly: which hospital currently has usable units of a given group, and what is the status of an outstanding request?")
body("This project designs and implements an Online Blood Bank and Donor Management System as a working web application. The prototype uses Python 3.10 and Flask 3.1 for the application layer, Jinja templates and a custom CSS design system for the presentation layer, and a normalized SQLite database for storage. The schema contains six related tables — users, hospitals, inventory, requests, donations and appointments — with primary keys, foreign keys, uniqueness rules, check constraints and search indexes. Two roles are supported. Donors register, sign in, check their 90-day eligibility, book a donation slot, search availability by blood group and city, raise requests and review their history. Hospital administrators manage hospitals, stock lines, requests, donations and donor records from a dispatch-board interface.")
body("Requests move through an explicit state machine: Pending → Approved → Completed, with cancellation permitted before fulfilment. The server checks available stock before marking a request complete and deducting units. Security controls include salted password hashing through Werkzeug, session-based authentication, role guards, CSRF protection on every POST, parameterized SQL and server-side validation. The interface was tested at desktop and mobile widths, including a reduced-motion path.")
body("Fourteen functional and security tests and five browser-level checks were executed. All executed checks passed: public pages returned 200, a missing CSRF token returned 400, a donor attempting an administrator page returned 403, an invalid status transition returned 409, completing a request reduced stock from 12 to 11 units, and an SQL-injection-style search returned zero rows. The result is a coherent academic prototype that demonstrates relational design, role-based workflows, validation and interface craft in a socially relevant domain.")
p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(6)
add_run(p,"Keywords: ",size=9.4,color=INK,bold=True)
add_run(p,"blood bank, relational database, Flask, SQLite, CRUD, role-based access, CSRF, normalisation, healthcare information systems.",size=9.4,color=SOFT,italic=True)
callout("Scope note","This is an academic prototype. It does not verify medical compatibility, screen donated blood, confirm donor eligibility, or replace hospital procedures. It must not be used for clinical decision-making.")
doc.add_page_break()

# ---------- contents ----------
doc.add_heading("Contents",level=1)
toc=[("1. Introduction","[[TOC1]]"),("2. System Analysis","[[TOC2]]"),("3. Requirements","[[TOC3]]"),
     ("4. System Design","[[TOC4]]"),("5. Database Design","[[TOC5]]"),("6. Implementation","[[TOC6]]"),
     ("7. Interface Walkthrough","[[TOC7]]"),("8. Testing and Validation","[[TOC8]]"),
     ("9. Results and Discussion","[[TOC9]]"),("10. Limitations and Future Scope","[[TOC10]]"),
     ("11. Conclusion","[[TOC11]]"),("References","[[TOC12]]"),("Appendix A — Run Instructions and Demo Accounts","[[TOC13]]")]
for label,token in toc:
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(7)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.55),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
    add_run(p,label,size=10.2,color=INK)
    add_run(p,"\t"+token,size=10.2,color=DEEP,bold=True)
callout("Implementation note","The original PBL brief specified MySQL and Node.js/Express. The submitted prototype uses Flask and SQLite so that the project runs on a single student laptop without an external database service. The relational schema, keys, joins and constraints are portable to MySQL 8; the change is an engineering decision, not a reduction in database scope.")
doc.add_page_break()

# ---------- 1 Introduction ----------
doc.add_heading("1. Introduction",level=1)
doc.add_heading("1.1 Background",level=2)
body("Blood banks sit inside a narrow window between supply and need. A unit has a shelf life, donors have a cooling-off period between donations, and a hospital request may be routine or immediately critical. When donor records, stock sheets and request logs live in separate places, staff spend time reconciling information instead of acting on it. A small, well-structured database application can reduce that friction by giving every desk the same current view.")
body("The project applies relational database principles to this scenario. It is deliberately scoped as a prototype: the goal is to demonstrate a correct schema, enforced relationships, role-based workflows, search and a usable interface, not to reproduce a regulated blood-bank information system.")
doc.add_heading("1.2 Problem Statement",level=2)
body("Manual or disconnected records create four recurring problems: duplicate donor entries, stale or inconsistent stock counts, requests whose status is known only to the person who took the call, and searches that depend on memory or a phone tree. The system addresses these by centralising donor, hospital, inventory, request and donation data in one relational store, with one workflow for each record type.")
doc.add_heading("1.3 Objectives",level=2)
for x in [
 "Design a normalized relational schema with six related entities and explicit integrity constraints.",
 "Implement donor registration, authentication and a personal dashboard with eligibility and history.",
 "Give hospital administrators a live board for stock, pending requests, donations and donor records.",
 "Enforce a request lifecycle in which status changes and stock deductions are validated on the server.",
 "Provide blood-group and city search that uses parameterised SQL joins rather than string concatenation.",
 "Protect state-changing operations with session authentication, role guards and CSRF tokens.",
 "Deliver a responsive, accessible interface that remains usable on a 390px-wide phone."]:
    bullet(x)
doc.add_heading("1.4 Scope",level=2)
body("In scope: donor registration and login; donor eligibility and donation history; appointment booking; public and authenticated availability search; donor-raised requests; administrator CRUD for hospitals, inventory, requests and donations; donor directory; role-based dashboards; JSON status updates; and a responsive design system.")
body("Out of scope: clinical compatibility testing, laboratory screening, donor identity verification, payment, SMS/email delivery, live multi-branch synchronisation, regulatory audit trails and production deployment hardening. These exclusions are treated again in Section 10.")

# ---------- 2 System Analysis ----------
doc.add_heading("2. System Analysis",level=1)
doc.add_heading("2.1 Limitations of the Existing Process",level=2)
for x in [
 ("Fragmented records. ","Donor details, stock registers and request books are maintained by different people and rarely share a key."),
 ("Duplicated donors. ","The same person can be registered more than once, which inflates counts and splits donation history."),
 ("Stale inventory. ","Stock is updated after the fact, so a search result may not reflect what is physically on the shelf."),
 ("Opaque requests. ","A requester has no reliable way to see whether a request is pending, approved or fulfilled."),
 ("Weak reporting. ","Aggregating units by group, hospital or city requires manual counting.")]:
    bullet(x[1],x[0])
doc.add_heading("2.2 Proposed System",level=2)
body("The proposed system is a single web application with two authenticated experiences and one public search surface. A donor owns a profile, donation history, appointments and requests. An administrator owns the hospital directory, stock sheet, request queue, donation log and donor directory. The public can search availability without signing in, but must authenticate to book, request or manage records.")
body("Every screen is generated from the same relational data. The board does not maintain a second copy of stock; it reads the inventory table. A request does not carry a hospital name; it stores a hospital_id and joins to the hospital record. This keeps the interface and the database consistent by construction.")
doc.add_heading("2.3 Users and Roles",level=2)
table(["Role","Primary goals","Permitted actions"],[
 ["Donor","Find blood, donate, track personal history","Register, sign in, search, book/cancel a slot, raise a request, view own donations and request status"],
 ["Hospital administrator","Keep stock accurate and clear requests","Manage hospitals and inventory, approve/decline/complete requests, log donations, view donor directory"],
 ["Public visitor","Check availability before travelling","Search stock by blood group and city; register or sign in"],
], [1.35,2.05,3.2])
caption("Table 2.1 — User roles and capabilities.")

# ---------- 3 Requirements ----------
doc.add_heading("3. Requirements",level=1)
doc.add_heading("3.1 Functional Requirements",level=2)
table(["ID","Requirement","Implementation evidence"],[
 ["FR-01","A visitor can register as a donor with name, email, password, phone, city and blood group.","Registration form; unique email; Werkzeug password hash; automatic sign-in."],
 ["FR-02","A registered user can sign in and is routed to the correct dashboard.","Session login; admin → /admin, donor → /dashboard."],
 ["FR-03","A donor can see the last donation, next eligible date, donation count and units given.","90-day cooldown calculation in donor_stats()."],
 ["FR-04","A donor can book or cancel a donation appointment.","Appointment table; future-date validation; ownership check on cancel."],
 ["FR-05","Any visitor can search stock by blood group and city.","Parameterized join across inventory and hospitals; results show hospital, city, phone and units."],
 ["FR-06","A donor can raise a request for a hospital, group, unit count and urgency.","Validation of group, 1–20 units and urgency; request stored as Pending."],
 ["FR-07","An administrator can add hospitals and maintain stock lines.","Hospital form creates eight inventory rows; stock form enforces 0–999 units."],
 ["FR-08","An administrator can approve, decline, complete or cancel requests.","Server-enforced state machine; stock check before completion; JSON API."],
 ["FR-09","An administrator can log a donation and update stock.","Donation insert; stock increment for the donor's group at the selected hospital."],
 ["FR-10","An administrator can review the donor directory and donation history.","Joined queries with last donation and total donation count."],
], [0.55,3.15,2.9], font=8.2, first_col_mono=True)
caption("Table 3.1 — Functional requirements traced to implemented behaviour.")

doc.add_heading("3.2 Non-Functional Requirements",level=2)
table(["ID","Quality","Target and evidence"],[
 ["NFR-01","Security","Salted password hashing; HttpOnly, SameSite=Lax session cookie; CSRF on all POSTs; role guards; parameterized SQL."],
 ["NFR-02","Data integrity","Primary and foreign keys, CHECK constraints, UNIQUE(hospital_id, blood_group), NOT NULL fields."],
 ["NFR-03","Usability","Clear navigation, explicit empty states, inline validation, visible request status and destructive-action confirmation copy."],
 ["NFR-04","Responsiveness","Desktop 1440px and mobile 390px layouts; no horizontal page overflow at either width."],
 ["NFR-05","Accessibility","WCAG AA contrast for text; 40px mobile buttons; keyboard focus rings; reduced-motion path."],
 ["NFR-06","Portability","Single Python process with a file-backed database; no external services required for the demo."],
 ["NFR-07","Maintainability","One schema file, one stylesheet, one script; reusable Jinja macros and icon set."],
], [0.6,1.25,4.75], font=8.3, first_col_mono=True)
caption("Table 3.2 — Non-functional requirements.")
doc.add_heading("3.3 Assumptions and Constraints",level=2)
for x in ["All records are synthetic and clearly labelled “Demo data” in the interface.",
          "The prototype is single-node and intended for one concurrent demonstration session.",
          "No medical, laboratory or regulatory validation is performed or implied.",
          "Passwords must be at least eight characters; the demo accounts use documented credentials.",
          "The local development server is used without HTTPS; production deployment would require TLS and hardened configuration."]:
    bullet(x)

# ---------- 4 System Design ----------
doc.add_heading("4. System Design",level=1)
doc.add_heading("4.1 Architecture",level=2)
body("The application follows a three-layer architecture. The presentation layer is server-rendered HTML with a small amount of JavaScript for clocks, counters, toasts, scroll reveals and asynchronous request-status updates. The application layer is a Flask process containing authentication, CSRF protection, role guards, validation and SQL queries. The data layer is SQLite, accessed through Python's sqlite3 module with a row factory and per-request connection.")
figure(ASSETS/"architecture.png","Figure 4.1 — Three-layer architecture. Browser requests pass through authentication, role checks and CSRF validation before parameterized SQL reaches SQLite.",6.45)
doc.add_heading("4.2 Module Decomposition",level=2)
table(["Module","Responsibility","Key routes"],[
 ["Public","Landing page and availability search","/ , /search"],
 ["Authentication","Registration, login, logout, session and CSRF","/login, /register, /logout"],
 ["Donor portal","Eligibility, history, appointments, requests","/dashboard, /book, /appointment/<id>/cancel, /request"],
 ["Admin board","Stock overview, pending request rail, recent donations","/admin"],
 ["Admin records","Hospitals, inventory, requests, donations, donors","/admin/hospitals, /admin/inventory, /admin/requests, /admin/donations, /admin/donors"],
 ["Status API","Validated status transitions and stock deduction","/api/requests/<id>/status"],
], [1.35,3.0,2.25], font=8.4)
caption("Table 4.1 — Application modules.")
doc.add_heading("4.3 Request Lifecycle",level=2)
body("A request begins as Pending and appears on the administrator's paper-slip rail. Approval changes the state to Approved but does not move stock: a request is still a request. Completion is the only transition that deducts units, and it is refused if the selected hospital does not hold enough units for that blood group. Cancellation is allowed from Pending or Approved; Completed and Cancelled are terminal.")
figure(ASSETS/"lifecycle.png","Figure 4.2 — Enforced request lifecycle. Completion performs the stock check and deduction; cancellation is terminal.",6.45)
doc.add_heading("4.4 Data Flow",level=2)
for i,x in enumerate([
 "A user opens a page or submits a form in the browser.",
 "Flask resolves the route and applies session, role and CSRF checks.",
 "The view validates input and runs parameterized SQL through the shared query helpers.",
 "SQLite returns rows or confirms a write; foreign keys and CHECK constraints reject invalid data.",
 "Jinja renders the updated page, or the JSON API returns a status payload to the browser.",
 "JavaScript updates a counter, stamp or toast without a full page reload where appropriate."],1):
    p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(0.28); p.paragraph_format.first_line_indent=Inches(-0.22); p.paragraph_format.space_after=Pt(3)
    add_run(p,f"{i}.  ",size=10.2,color=RED,bold=True); add_run(p,x,size=10.2)

# ---------- 5 Database ----------
doc.add_heading("5. Database Design",level=1)
doc.add_heading("5.1 Entity–Relationship Model",level=2)
body("The schema contains six entities. Users own requests and donations and may book appointments. Hospitals hold inventory lines, receive requests, receive donations and host appointments. Inventory is unique per hospital and blood group, which prevents duplicate stock rows and makes an upsert-style update safe. Requests and donations reference both a user and a hospital, preserving the history needed to explain a stock movement.")
figure(ASSETS/"er.png","Figure 5.1 — Entity–relationship model with primary keys, foreign keys and one-to-many relationships.",6.55)
doc.add_heading("5.2 Tables and Relationships",level=2)
table(["Table","Purpose","Key fields and relationships"],[
 ["users","One row per donor or administrator","user_id PK; email UNIQUE; role CHECK (donor/admin); blood_group"],
 ["hospitals","Hospital and blood-centre directory","hospital_id PK; name; city; phone; address"],
 ["inventory","Stock per hospital and group","inventory_id PK; hospital_id FK → hospitals; UNIQUE(hospital_id, blood_group); units CHECK ≥ 0"],
 ["requests","Blood requests raised by donors or staff","request_id PK; requester_id FK → users; hospital_id FK → hospitals; status CHECK"],
 ["donations","Recorded donation events","donation_id PK; donor_id FK → users; hospital_id FK → hospitals; units; date"],
 ["appointments","Booked donation slots","appointment_id PK; donor_id FK; hospital_id FK; slot_date; slot_time; status"],
], [1.05,2.0,3.55], font=8.2, first_col_mono=True)
caption("Table 5.1 — Entity overview and key relationships.")
doc.add_heading("5.3 Constraints and Indexes",level=2)
table(["Constraint","Purpose"],[
 ["PRIMARY KEY","Uniquely identifies every record and gives foreign keys a stable target."],
 ["FOREIGN KEY","Prevents inventory, requests, donations or appointments from referencing a missing user or hospital."],
 ["UNIQUE(email)","Prevents duplicate donor accounts."],
 ["UNIQUE(hospital_id, blood_group)","Guarantees one stock line per group per hospital."],
 ["CHECK(units ≥ 0) / CHECK(units > 0)","Rejects negative stock and non-positive request or donation quantities."],
 ["CHECK(role) and CHECK(status)","Restricts roles and workflow states to known values."],
 ["ON DELETE CASCADE (inventory)","Removes stock lines only when their parent hospital is deleted; request history blocks deletion first."],
 ["idx_inventory_group, idx_requests_status, idx_donations_donor","Speed up the most frequent filters and joins."],
], [2.15,4.4], font=8.3, first_col_mono=True)
caption("Table 5.2 — Integrity rules and indexes.")
doc.add_heading("5.4 Normalisation",level=2)
body("The schema is in Third Normal Form. Each field is atomic (1NF). Stock, requests and donations use whole-row primary keys, so no non-key attribute depends on only part of a candidate key (2NF). Descriptive data is stored once: a request stores hospital_id rather than hospital name or city, and the display joins to hospitals. A donor's blood group lives in users, not repeated in every donation row. Removing any of these indirect dependencies would lose meaning, so no transitive dependency remains (3NF).")
body("Inventory is the one intentional exception to a pure event log: it stores the current count, while donations and requests provide the history that explains changes. This is a deliberate read-performance choice for a dashboard, and the trade-off is documented rather than hidden.")
doc.add_heading("5.5 History and Deletion",level=2)
body("Healthcare records should not disappear casually. Hospitals with requests cannot be deleted; the administrator receives an explanatory message. Inventory rows cascade only when their parent hospital is legitimately removed. Requests and donations are retained in their terminal states, so the board can still explain why stock changed. A production system would add soft deletion, audit tables and role-specific retention rules.")

# ---------- 6 Implementation ----------
doc.add_heading("6. Implementation",level=1)
doc.add_heading("6.1 Technology Stack",level=2)
table(["Layer","Technology","Why it was chosen"],[
 ["Language","Python 3.10","Readable server code and a mature standard library."],
 ["Web framework","Flask 3.1.3","Small surface area; explicit routes; easy to demonstrate in a PBL setting."],
 ["Templates","Jinja2","Server-rendered HTML with reusable macros and autoescaping."],
 ["Database","SQLite 3","Zero-configuration file database; supports foreign keys, CHECK constraints and joins."],
 ["Front end","HTML5, CSS3, vanilla JavaScript","No build step; full control over the dispatch-board visual language."],
 ["Security","Werkzeug password hashing, signed sessions, CSRF tokens","Standard, well-understood primitives without custom cryptography."],
 ["Testing","Flask test client, Playwright (Chrome)","Route-level assertions and real-browser layout/contrast checks."],
], [1.15,2.15,3.25], font=8.35)
caption("Table 6.1 — Implementation stack.")
callout("Deviation from the original brief","The brief named MySQL 8 and Node.js/Express. The implemented prototype uses SQLite and Flask so it can run from a single folder on a student machine. The data model, keys, constraints and join queries map directly to MySQL; Section 10 lists the migration path.")
doc.add_heading("6.2 Project Structure",level=2)
code_block(["blood-bank-app/","├── app.py                 # schema, seed, routes, auth, API","├── templates/             # Jinja pages and reusable macros","├── static/css/style.css   # design system and motion","├── static/js/app.js       # clock, counters, stamps, toasts","├── static/fonts/          # local Archivo and JetBrains Mono","└── instance/bloodbank.db  # SQLite database (created on first run)"])
doc.add_heading("6.3 Authentication, Roles and CSRF",level=2)
body("Registration validates the email, password length, blood group and required fields, hashes the password with Werkzeug and stores the new user with the donor role. Login verifies the hash and writes user_id and role into the signed session. Two decorators protect the application: login_required redirects anonymous visitors to the sign-in page, while admin_required returns 403 for a signed-in donor who tries to open an administrator route.")
body("Every POST is checked by a before_request hook. The token is generated per session, embedded in forms and in a meta tag for fetch calls, and compared using secrets.compare_digest. The API accepts the token through the X-CSRF-Token header; a missing or stale token returns HTTP 400 with a JSON error.")
code_block(["@app.before_request","def csrf_protect():","    if request.method == \"POST\":","        sent = request.form.get(\"csrf_token\") or request.headers.get(\"X-CSRF-Token\")","        if not sent or not secrets.compare_digest(sent, session.get(\"csrf\", \"\")):","            if request.path.startswith(\"/api/\"):","                return jsonify({\"error\": \"Invalid CSRF token.\"}), 400","            abort(400)"])
doc.add_heading("6.4 Request State Machine and Stock Check",level=2)
body("Status changes are not accepted as free text. The API defines the allowed next states for each current state and returns HTTP 409 for anything else. Completion performs a stock read for the request's hospital and group; if the row is missing or holds fewer units than requested, the transition is refused. Only then does the inventory row decrement and the request become Completed. The browser receives the new status and the updated total for that blood group, which drives the counter animation.")
code_block(["allowed = {","    \"Pending\":   {\"Approved\", \"Cancelled\"},","    \"Approved\":  {\"Completed\", \"Cancelled\"},","    \"Completed\": set(),","    \"Cancelled\": set(),","}","# ...","if new_status == \"Completed\":","    inv = get_inventory(row[\"hospital_id\"], row[\"blood_group\"])","    if not inv or inv[\"units\"] < row[\"units\"]:","        return jsonify({\"error\": \"Not enough units in stock.\"}), 409","    deduct_units(inv[\"inventory_id\"], row[\"units\"])"])
doc.add_heading("6.5 Search and Query Design",level=2)
body("Search is built from optional predicates rather than one fixed query. A blood group uses exact equality; a city uses a parameterized LIKE. The query joins inventory to hospitals so a result can show the site name, city and phone without duplicating those fields in the inventory table. Only rows with positive stock are returned, ordered by available units and hospital name.")
code_block(["SELECT i.blood_group, i.units, h.name, h.city, h.phone","FROM inventory i","JOIN hospitals h ON h.hospital_id = i.hospital_id","WHERE i.blood_group = ?","  AND h.city LIKE ?","  AND i.units > 0","ORDER BY i.units DESC, h.name;"])
body("Values are always bound as parameters. A search for a quote-based injection payload is treated as a literal city string and returns no rows, as recorded in Table 8.2.")
doc.add_heading("6.6 Interface and Motion",level=2)
body("The interface is built around a dispatch-board metaphor rather than a generic admin template: cream paper, an oxidized-red rail, ink typography, amber signals and teal approvals. Archivo carries the interface; JetBrains Mono carries references, dates and stock counts. Stock is shown as paper strips with split-flap counters; pending requests are paper slips; approval presses a rubber stamp and then lifts the slip away.")
body("Motion is purposeful and short: a 260ms counter tick, a 620ms stamp, a 320ms toast and a 60ms stagger on scroll reveals. The reduced-motion media query collapses animation and transition durations and reveals all sections immediately. Nothing relies on animation to communicate state; the status text and stock value are always present.")
doc.add_heading("6.7 Validation and Error Handling",level=2)
for x in ["Required fields, email format, password length and blood-group membership are checked server-side.",
          "Units are bounded to 1–20 for requests and 0–999 for stock; appointments reject past dates.",
          "Ownership checks prevent a donor from cancelling another donor's appointment.",
          "403 and 404 pages use the same visual world and offer a route back to the board.",
          "Flash messages confirm successful writes and explain refused operations in plain language."]:
    bullet(x)

# ---------- 7 Interface ----------
doc.add_heading("7. Interface Walkthrough",level=1)
body("The following screenshots were captured from the running application with the seeded demo database. They are evidence of the rendered interface, not mock-ups.")
doc.add_heading("7.1 Public Landing and Availability Search",level=2)
body("The landing page states the problem and offers two immediate actions: search availability or register as a donor. The search form accepts a blood group and a city; the result row gives the hospital, city, phone number and units, so a visitor can act without signing in. The footer exposes group-level totals and flags low groups.")
figure(ASSETS/"landing.jpg","Figure 7.1 — Public landing page at 1440×900. The board metaphor is established before any data is shown.",6.35)
figure(ASSETS/"search.jpg","Figure 7.2 — Public availability search for O+ in Pune, returning City Hospital with its contact number.",6.35)
doc.add_heading("7.2 Administrator Dispatch Board",level=2)
body("The board is the administrator's first viewport. Group stock is arranged as split-flap strips, with critical groups marked in red. Pending requests sit in a paper-slip rail; each slip carries the reference, group, units, hospital, urgency and the two decisions. The right column shows recent donations and hospital stock totals. The top statistics summarise units, critical groups, pending requests and registered donors.")
figure(ASSETS/"board.jpg","Figure 7.3 — Administrator dispatch board at 1440×900, showing stock strips, pending request slips and recent donations.",6.4)
doc.add_heading("7.3 Request Rail and Status Actions",level=2)
body("The request page exposes the full queue with status filters. Pending requests offer Approve and Decline; approved requests offer Complete and Cancel. Completing a request is the only action that changes stock, and the confirmation is shown as a toast while the stock counter animates. Terminal rows are retained with a Filed label.")
figure(ASSETS/"requests.jpg","Figure 7.4 — Request queue with urgency, status, hospital and permitted actions.",6.3)
doc.add_heading("7.4 Donor Dashboard",level=2)
body("The donor dashboard answers the questions a donor actually asks: when did I last donate, when can I donate again, what have I given, and what is happening to my request? Eligibility is calculated from the most recent donation plus the 90-day cooldown. The same page carries appointment booking, a request form and personal request history.")
figure(ASSETS/"donor.jpg","Figure 7.5 — Donor dashboard with eligibility strip, donation history, appointment booking and request history.",5.85)
doc.add_heading("7.5 Responsive Behaviour",level=2)
body("At 390px the rail becomes a compact header with a horizontally scrollable tab strip, the board collapses to one column, stock strips become two-up, and action buttons grow to a 40px minimum height. The mobile layout has no horizontal page overflow; wide tables scroll inside their own container. The active tab is automatically centred on load.")
figure(ASSETS/"mobile-pair.jpg","Figure 7.6 — Mobile layout at 390×844: administrator board (left) and public landing page (right).",5.35)

# ---------- 8 Testing ----------
doc.add_heading("8. Testing and Validation",level=1)
doc.add_heading("8.1 Test Environment",level=2)
table(["Item","Value"],[
 ["Application","Flask 3.1.3, Python 3.10.18, SQLite 3"],
 ["Database state","Seeded demo database: 1 administrator, 5 donors, 3 hospitals, 24 stock lines, 136 units, 6 requests, 7 donations, 1 appointment"],
 ["Route tests","Flask test client with a fresh seeded database"],
 ["Browser tests","Google Chrome via Playwright at 1440×900 and 390×844"],
 ["Checks","HTTP status, database row counts, stock before/after, transition errors, horizontal overflow, contrast, reduced motion"],
], [1.35,5.2], font=8.5)
caption("Table 8.1 — Test environment and starting data.")
doc.add_heading("8.2 Functional and Security Tests",level=2)
table(["ID","Test","Expected","Actual","Result"],[
 ["T01","Open public landing page","200 and rendered board","200","Pass"],
 ["T02","Search O+ in Pune","Matching hospital rows only","200; City Hospital returned","Pass"],
 ["T03","POST login without CSRF token","Rejected","400","Pass"],
 ["T04","Admin login","Redirect to /admin","302 → /admin","Pass"],
 ["T05","Donor login","Redirect to /dashboard","302 → /dashboard","Pass"],
 ["T06","Donor opens /admin","Forbidden","403","Pass"],
 ["T07","Donor calls status API","Forbidden","403","Pass"],
 ["T08","Approve a pending request","Status Approved; stock unchanged","200; Approved","Pass"],
 ["T09","Complete an approved request","Status Completed; stock deducted","200; 12 → 11 units","Pass"],
 ["T10","Repeat a completed transition","Rejected by state machine","409","Pass"],
 ["T11","Raise a valid donor request","New Pending row","302; count 6 → 7","Pass"],
 ["T12","Book a future appointment","New Booked row","302; count 1 → 2","Pass"],
 ["T13","SQL-injection-style city search","Treated as data; no leak","200; 0 result rows","Pass"],
 ["T14","Cancel another donor's appointment","Not found / refused","404","Pass (ownership check)"],
], [0.46,1.98,2.08,1.46,0.62], font=7.9, first_col_mono=True)
caption("Table 8.2 — Executed functional and security tests. The database was reset after the test run.")
doc.add_heading("8.3 Browser and Accessibility Checks",level=2)
table(["Check","Method","Result"],[
 ["Horizontal overflow at 1440px","document.scrollWidth vs clientWidth","No overflow"],
 ["Horizontal overflow at 390px","document.scrollWidth vs clientWidth","No page overflow; tables scroll internally"],
 ["Text contrast","Computed foreground/background luminance","All sampled text meets WCAG AA"],
 ["Reduced motion","prefers-reduced-motion: reduce","Counters show final values; reveals instant"],
 ["Loaded assets","Chrome console and network","No application errors; local fonts loaded"],
], [1.9,3.0,1.65], font=8.2)
caption("Table 8.3 — Browser-level checks.")
doc.add_heading("8.4 Discussion of Results",level=2)
body("The executed tests cover the paths most likely to fail in a role-based CRUD application: authentication, authorisation, CSRF, state transitions, stock consistency and search safety. The most important result is not that a page rendered, but that the server refused invalid actions. A donor could not reach an administrator page, a completed request could not be approved again, and stock changed only when a valid completion was processed.")
body("The browser checks address a different risk: a desktop-only interface that looks correct in a screenshot but breaks on a phone or excludes a keyboard user. The mobile navigation, table wrappers and minimum button sizes were adjusted after finding a 20px overflow and several sub-38px controls; the final build reports no page-level overflow and passes sampled contrast checks. The test set is deliberately small and transparent, not an exhaustive penetration test or clinical validation.")

# ---------- 9 Results ----------
doc.add_heading("9. Results and Discussion",level=1)
body("The prototype meets the seven objectives stated in Section 1.3. The schema is normalized, the relationships are enforced by the database, and the application exposes a coherent workflow for each role. The result is best read as a chain of evidence: schema → route → validation → rendered state → recorded test.")
table(["Objective","Evidence"],[
 ["Relational schema","Six tables with primary keys, foreign keys, CHECK and UNIQUE constraints; ER model in Figure 5.1."],
 ["Donor portal","Registration and login; eligibility, history, appointment booking and personal requests on one dashboard."],
 ["Administrator board","Stock strips, pending slips, recent donations, hospital totals and request queue."],
 ["Server-enforced lifecycle","Allowed transitions table; 409 on invalid transition; stock check before deduction."],
 ["Parameterized search","Join query with bound parameters; injection-style input returned zero rows."],
 ["Security controls","Password hashing; role guards return 403; missing CSRF token returns 400."],
 ["Responsive, accessible UI","390px layout without page overflow; WCAG AA contrast; reduced-motion support."],
], [1.65,4.9], font=8.4)
caption("Table 9.1 — Objectives mapped to implementation evidence.")
body("Two results are worth highlighting. First, the stock completion test produced a measurable state change (12 → 11 units) rather than a cosmetic success message, which demonstrates that the interface and the database are connected through the validated API. Second, the authorisation tests returned 403 rather than a redirect to a generic page, which shows that the role check is enforced at the route boundary.")
body("The design also contributes to the result. A conventional dashboard would have displayed the same data, but the dispatch-board language gives the administrator a mental model that matches the work: stock is a strip, a request is a slip, approval is a stamp, and completion is a deduction. That metaphor is not decoration; it makes the state of the system legible at a glance.")

# ---------- 10 Limitations ----------
doc.add_heading("10. Limitations and Future Scope",level=1)
doc.add_heading("10.1 Limitations",level=2)
for x in [
 "SQLite is a single-node database. It is adequate for a demonstration but not for concurrent multi-branch operation.",
 "The prototype does not verify clinical compatibility, screen blood, confirm donor identity or enforce medical eligibility.",
 "There is no email/SMS notification, audit trail, rate limiting or production-grade backup strategy.",
 "Search is not paginated; a large inventory would need limits, total counts and cursor-based navigation.",
 "The local server runs over HTTP; deployment would require TLS, secure cookies, a production WSGI server and secret management.",
 "Automated tests cover core workflows but are not a substitute for a full accessibility audit or penetration test."]:
    bullet(x)
doc.add_heading("10.2 Future Scope",level=2)
for x in [
 "Migrate the schema to MySQL 8 or PostgreSQL and add versioned migrations for multi-user deployment.",
 "Add audit logging and soft deletion for donor, inventory and request changes.",
 "Introduce notification workflows for request status, appointment reminders and low-stock alerts.",
 "Add pagination, saved searches and CSV export for administrators.",
 "Integrate a map or geo-search for nearby hospitals, with consent and privacy controls.",
 "Expand automated tests into a CI pipeline, including accessibility checks and load testing.",
 "Add configurable inventory thresholds and expiry tracking for blood units."]:
    bullet(x)

# ---------- 11 Conclusion ----------
doc.add_heading("11. Conclusion",level=1)
body("This project turns a relational database exercise into a working healthcare workflow. The six-table schema captures donors, hospitals, stock, requests, donations and appointments with explicit keys and constraints. The Flask application gives each role a focused surface: donors can find blood and manage their own contribution history; administrators can see the shelf, clear the request rail and record donations. The request state machine ensures that stock moves only after a valid, checked completion.")
body("The implementation also demonstrates that database quality is visible in the interface. When relationships are enforced, a hospital name cannot drift between screens; when stock is read from one table, the counter and the completion result agree; when CSRF and role guards are present, an invalid action is refused rather than silently ignored. These are the behaviours the tests exercised.")
body("The prototype is not a clinical product. It is an academic demonstration of schema design, server-side validation, role-based access and interface craft in a domain where those qualities matter. With the migration, audit and notification work described in Section 10, the same structure could serve as the foundation for a larger blood-bank information system.")

# ---------- References ----------
doc.add_page_break()
doc.add_heading("References",level=1)
refs=[
 "Elmasri, R. and Navathe, S. B., Fundamentals of Database Systems, 7th ed. Pearson, 2017.",
 "Pallets Projects, Flask Documentation (3.1.x), 2026. [Online]. Available: https://flask.palletsprojects.com/",
 "SQLite Consortium, SQLite Documentation, 2026. [Online]. Available: https://www.sqlite.org/docs.html",
 "Mozilla Foundation, MDN Web Docs: HTML, CSS and JavaScript, 2026. [Online]. Available: https://developer.mozilla.org/",
 "OWASP Foundation, Cross-Site Request Forgery Prevention Cheat Sheet, 2025. [Online]. Available: https://cheatsheetseries.owasp.org/",
 "W3C, Web Content Accessibility Guidelines (WCAG) 2.2, 2023. [Online]. Available: https://www.w3.org/TR/WCAG22/",
 "World Health Organization, Blood safety and availability, 2024. [Online]. Available: https://www.who.int/news-room/fact-sheets/detail/blood-safety-and-availability",
 "Werkzeug Documentation, Password Hashing, 2026. [Online]. Available: https://werkzeug.palletsprojects.com/",
]
for i,r in enumerate(refs,1):
    p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(0.3); p.paragraph_format.first_line_indent=Inches(-0.3)
    p.paragraph_format.space_after=Pt(6)
    add_run(p,f"[{i}]  ",size=9.2,color=RED,bold=True); add_run(p,r,size=9.2,color=SOFT)

# ---------- Appendix ----------
doc.add_page_break()
doc.add_heading("Appendix A — Run Instructions and Demo Accounts",level=1)
doc.add_heading("A.1 Local Setup",level=2)
code_block(["cd blood-bank-app","python3 -m venv .venv","source .venv/bin/activate","pip install -r requirements.txt","python app.py","# open http://127.0.0.1:5000"])
body("The database is created and seeded automatically on first run in instance/bloodbank.db. To reset the demonstration, stop the server, delete that file and start the server again.")
doc.add_heading("A.2 Demo Accounts",level=2)
table(["Role","Email","Password"],[
 ["Administrator","admin@bloodbank.test","Admin@123"],
 ["Donor","donor@bloodbank.test","Donor@123"],
 ["Donor (secondary)","sneha@bloodbank.test","Donor@123"],
], [1.45,2.55,2.55], font=8.7)
caption("Table A.1 — Demo credentials. All accounts and records are synthetic.")
doc.add_heading("A.3 Data Reset",level=2)
body("Deleting instance/bloodbank.db and restarting the application recreates the schema and reseeds the demonstration. No production data is used at any point.")
p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(18)
callout("Academic integrity","The implementation, screenshots and test results in this report describe the submitted prototype. Where the implementation differs from the original design brief — notably Flask and SQLite instead of Express and MySQL — the deviation is stated explicitly in Section 6.1 rather than hidden.")

# ---------- final metadata ----------
doc.core_properties.title="Online Blood Bank and Donor Management System — PBL Report"
doc.core_properties.subject="Project-Based Learning implementation report"
doc.core_properties.author="[Your Name]"
doc.core_properties.keywords="blood bank, Flask, SQLite, relational database, PBL"
doc.save(DOCX)
print("docx saved",DOCX, os.path.getsize(DOCX))

# ---------- PDF conversion + TOC page numbers ----------
def convert():
    subprocess.run(["soffice","--headless","--convert-to","pdf","--outdir",str(BASE),str(DOCX)],check=True,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
    return BASE/"Blood_Bank_Donor_Management_Report.pdf"

convert()
from pypdf import PdfReader
heads=["1. Introduction","2. System Analysis","3. Requirements","4. System Design","5. Database Design","6. Implementation","7. Interface Walkthrough","8. Testing and Validation","9. Results and Discussion","10. Limitations and Future Scope","11. Conclusion","References","Appendix A"]
pages={}
reader=PdfReader(PDF)
for i,page in enumerate(reader.pages):
    if i < 3:
        continue
    lines=[ln.strip() for ln in (page.extract_text() or "").splitlines() if ln.strip()]
    start=0
    for j,ln in enumerate(lines):
        if re.fullmatch(r"Page \d+ of \d+", ln):
            start=j+1
            break
    for ln in lines[start:]:
        for h in heads:
            if h in pages:
                continue
            if ln == h or ln.startswith(h + " "):
                pages[h]=i+1
# update tokens in docx
from docx import Document as D2
d=D2(DOCX)
for p in d.paragraphs:
    for h in heads:
        token="[[TOC%d]]"%(heads.index(h)+1)
        if token in p.text:
            for r in p.runs:
                if token in r.text:
                    r.text=r.text.replace(token,str(pages.get(h,"—")))
# also cover author if wanted left placeholder
d.save(DOCX)
convert()
print("pages",pages)
print("pdf saved",PDF,os.path.getsize(PDF))
